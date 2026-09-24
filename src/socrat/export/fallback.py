"""Small fixture-mode exporter until the full Word exporter is available."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument

from socrat.contracts import DiagnosticWork, ErrorAnalysisResult, student_view


def _document_bytes(document: DocxDocument) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FixtureExporter:
    def student_docx(self, work: DiagnosticWork) -> bytes:
        view = student_view(work)
        document = Document()
        document.add_heading(view.title, 0)
        for index, variant in enumerate(view.variants):
            if index:
                document.add_page_break()
            document.add_heading(variant.label, 1)
            document.add_paragraph(variant.instruction)
            for task in variant.tasks:
                document.add_heading(f"Задание {task.number}", 2)
                document.add_paragraph(task.text)
                if task.support:
                    document.add_paragraph(task.support)
                for _ in range(task.answer_lines):
                    document.add_paragraph("________________________________________")
            document.add_heading("Подумай и ответь", 2)
            for question in variant.reflection_questions:
                document.add_paragraph(question)
                document.add_paragraph("________________________________________")
        return _document_bytes(document)

    def teacher_docx(self, work: DiagnosticWork) -> bytes:
        document = Document()
        document.add_heading("Методические материалы для учителя", 0)
        document.add_paragraph(work.title)
        for variant in work.variants:
            document.add_heading(variant.student_label, 1)
            for task in variant.tasks:
                document.add_heading(f"Задание {task.number}", 2)
                document.add_paragraph(task.student_text)
                document.add_paragraph(f"Ответ: {task.expected_answer}")
        return _document_bytes(document)

    def analysis_docx(self, result: ErrorAnalysisResult) -> bytes:
        document = Document()
        document.add_heading("Анализ типичной ошибки", 0)
        document.add_paragraph(result.summary)
        return _document_bytes(document)

    def filenames(self, work: DiagnosticWork) -> tuple[str, str]:
        suffix = f"{work.request.grade}кл_{work.request.subject_name}.docx"
        return f"Задания_{suffix}", f"Методичка_{suffix}"
