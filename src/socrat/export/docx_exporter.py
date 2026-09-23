"""Word documents for students, teachers and error-analysis reports."""

from __future__ import annotations

import json
import re
from functools import lru_cache
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from socrat.config import get_settings
from socrat.contracts import (
    ERROR_LAYER_LABELS,
    LABELS_RU,
    DiagnosticWork,
    ErrorAnalysisResult,
    student_view,
)

ROOT = Path(__file__).resolve().parents[3]
TECHNIQUES_FILE = ROOT / "data" / "techniques.json"
GRAY = RGBColor(90, 90, 90)


def _configure(document: DocxDocument, *, landscape: bool = False) -> None:
    section = document.sections[0]
    section.page_width = Mm(297 if landscape else 210)
    section.page_height = Mm(210 if landscape else 297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)


def _bytes(document: DocxDocument) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _line(document: DocxDocument, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(value)


def _bullets(document: DocxDocument, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def _table(
    document: DocxDocument,
    headers: list[str],
    rows: list[list[str]],
    widths_mm: list[int] | None = None,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths_mm is None
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        if widths_mm:
            cell.width = Mm(widths_mm[index])
        for run in cell.paragraphs[0].runs:
            run.bold = True
    _repeat_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths_mm:
                cells[index].width = Mm(widths_mm[index])
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)
    document.add_paragraph()


@lru_cache(maxsize=1)
def _technique_names() -> dict[str, str]:
    data = json.loads(TECHNIQUES_FILE.read_text(encoding="utf-8"))
    return {item["technique_id"]: item["name"] for item in data["techniques"]}


def _observation_rules() -> dict:
    path = get_settings().case_reference_dir / "observation_rules.json"
    return json.loads(path.read_text(encoding="utf-8"))


def _safe_filename_part(value: str, maximum: int) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", value)
    cleaned = " ".join(cleaned.split()).strip(" .")
    return cleaned[:maximum].rstrip(" .") or "Тема"


class DocxExporter:
    def student_docx(self, work: DiagnosticWork) -> bytes:
        view = student_view(work)
        document = Document()
        _configure(document)
        for index, variant in enumerate(view.variants):
            if index:
                document.add_page_break()
            document.add_paragraph("Имя ______________   Класс ______   Дата __________")
            title = document.add_paragraph(style="Title")
            title.add_run(view.title).font.size = Pt(14)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_paragraph(f"{view.subject_name}, {view.grade} класс · {variant.label}")
            instruction = document.add_paragraph()
            instruction.add_run(variant.instruction).italic = True

            for task in variant.tasks:
                document.add_heading(f"Задание {task.number}", level=2)
                document.add_paragraph(task.text)
                if task.support:
                    support = task.support
                    if not support.lower().startswith("подсказка:"):
                        support = f"Подсказка: {support}"
                    run = document.add_paragraph().add_run(support)
                    run.italic = True
                    run.font.color.rgb = GRAY
                for _ in range(task.answer_lines):
                    document.add_paragraph("________________________________________")

            document.add_heading("Подумай и ответь", level=2)
            for question in variant.reflection_questions:
                document.add_paragraph(question, style="List Bullet")
                document.add_paragraph("________________________________________")
                document.add_paragraph("________________________________________")
        return _bytes(document)

    def filenames(self, work: DiagnosticWork) -> tuple[str, str]:
        subject = _safe_filename_part(work.request.subject_name, 40)
        topic = _safe_filename_part(work.request.topic, 40)
        suffix = f"{work.request.grade}кл_{subject}_{topic}.docx"
        return f"Задания_{suffix}", f"Методичка_{suffix}"

    def teacher_docx(self, work: DiagnosticWork) -> bytes:
        document = Document()
        _configure(document, landscape=True)
        document.add_paragraph("Методические материалы для учителя", style="Title")
        _line(document, "Тема", work.title)
        _line(document, "Предмет и класс", f"{work.request.subject_name}, {work.request.grade} класс")
        _line(document, "Дата", work.created_at.strftime("%d.%m.%Y"))
        _line(document, "Уровень", LABELS_RU[work.request.level.value])
        _line(document, "Фокус УУД", LABELS_RU[work.request.uud_focus.value])
        _line(document, "Число заданий", str(work.request.task_count))

        document.add_heading("Ограничения", level=1)
        box = document.add_table(rows=1, cols=1)
        box.style = "Table Grid"
        cell = box.cell(0, 0)
        cell.text = work.limitations[0] if work.limitations else ""
        if work.limitations:
            cell.paragraphs[0].style = "List Bullet"
        for limitation in work.limitations[1:]:
            cell.add_paragraph(limitation, style="List Bullet")

        for variant in work.variants:
            document.add_heading(
                f"{variant.student_label} — {LABELS_RU[variant.level.value].lower()} уровень", level=1
            )
            document.add_paragraph(variant.level_rationale)
            document.add_heading("План времени", level=2)
            plan = variant.time_plan
            _table(
                document,
                ["Этап", "Минуты"],
                [
                    ["Чтение", f"{plan.reading_minutes:g}"],
                    ["Задания", f"{plan.tasks_minutes:g}"],
                    ["Рефлексия", f"{plan.reflection_minutes:g}"],
                    ["Итого", f"{plan.total_minutes:g}"],
                ],
            )
            document.add_paragraph(plan.basis)
            for task in variant.tasks:
                document.add_heading(f"Задание {task.number}", level=2)
                document.add_paragraph(task.student_text)
                if task.support:
                    _line(document, "Опора", task.support)
                _line(document, "Предметная цель", task.subject_goal)
                _line(document, "Метапредметная цель", task.meta_goal)
                _line(document, "Ответ", task.expected_answer)
                if task.solution_steps:
                    document.add_heading("Решение по шагам", level=3)
                    for index, step in enumerate(task.solution_steps, start=1):
                        document.add_paragraph(f"{index}. {step.text}")
                if task.alternative_solutions:
                    document.add_heading("Допустимые другие способы", level=3)
                    _bullets(document, task.alternative_solutions)

                document.add_heading("Карта результатов", level=3)
                _table(
                    document,
                    [
                        "Тип",
                        "ID (внутр.)",
                        "Формулировка",
                        "Источник",
                        "Раздел",
                        "Стр.",
                        "Почему соответствует",
                        "✓ проверено",
                    ],
                    [
                        [
                            LABELS_RU[link.type.value],
                            link.outcome_id,
                            link.text,
                            link.source_title,
                            link.section,
                            str(link.page),
                            link.rationale,
                            "✓" if link.verified else (link.verification_note or "нет"),
                        ]
                        for link in task.outcome_links
                    ],
                    [22, 23, 54, 38, 31, 12, 57, 20],
                )

                document.add_heading("Наблюдаемые признаки УУД", level=3)
                _table(
                    document,
                    [
                        "Группа",
                        "Действие",
                        "Основание в задании",
                        "Что считать проявлением",
                        "Подсказки по шкале",
                    ],
                    [
                        [
                            LABELS_RU[indicator.group.value],
                            indicator.action,
                            indicator.trigger,
                            indicator.evidence,
                            "\n".join(
                                f"{LABELS_RU.get(status.value, status.value)}: {hint}"
                                for status, hint in indicator.scale_hints.items()
                            ),
                        ]
                        for indicator in task.uud_indicators
                    ],
                    [34, 51, 55, 55, 62],
                )

                if task.personal_orientation:
                    _line(
                        document,
                        "Личностная направленность",
                        f"{task.personal_orientation} Индивидуальный балл не ставится.",
                    )
                if task.typical_errors:
                    document.add_heading("Типичные ошибки", level=3)
                    _table(
                        document,
                        ["Как выглядит", "Слой", "Возможное объяснение", "Что сделать"],
                        [
                            [
                                error.description,
                                ERROR_LAYER_LABELS[error.layer],
                                error.interpretation,
                                error.teacher_move,
                            ]
                            for error in task.typical_errors
                        ],
                        [64, 42, 76, 75],
                    )
                if task.oral_questions:
                    document.add_heading("Устные вопросы", level=3)
                    _bullets(document, task.oral_questions)
                if task.technique_ids:
                    document.add_heading("Приёмы", level=3)
                    names = _technique_names()
                    _bullets(
                        document, [names.get(identifier, identifier) for identifier in task.technique_ids]
                    )
                if task.conducting_note:
                    _line(document, "Как проводить", task.conducting_note)
                document.add_heading("Автопроверки", level=3)
                checks = task.checks
                check_lines = [
                    "Арифметика проверена кодом ✓"
                    if checks.math_verified
                    else "Арифметика не проверялась"
                    if checks.math_verified is None
                    else "Арифметика не прошла проверку",
                    f"Ответ не раскрыт в ученическом листе: {'✓' if checks.answer_leak_free else 'нет'}",
                    f"УУД имеют основание: {'✓' if checks.uud_grounded else 'нет'}",
                    f"Ссылки на результаты проверены: {'✓' if checks.outcomes_verified else 'нет'}",
                    f"Уровень согласован: {'✓' if checks.level_consistent else 'нет'}",
                ]
                _bullets(document, check_lines + checks.notes)

            document.add_heading("Рефлексия", level=2)
            _bullets(document, variant.reflection.questions)
            if variant.reflection.reading_guide:
                _line(document, "Как читать ответы", variant.reflection.reading_guide)
            document.add_paragraph(variant.reflection.note)

        rules = _observation_rules()
        document.add_heading("Шкала наблюдений", level=1)
        document.add_paragraph(" / ".join(rules["uud"]))
        _bullets(document, rules["rules"])
        document.add_heading("Источники", level=1)
        for source in work.sources:
            document.add_paragraph(
                f"{source.title}. {source.url}" + (f" {source.note}" if source.note else "")
            )
        document.add_paragraph(
            "Коды результатов — внутренние идентификаторы системы, не номера пунктов ФГОС."
        )
        return _bytes(document)

    def analysis_docx(self, result: ErrorAnalysisResult) -> bytes:
        document = Document()
        _configure(document)
        document.add_paragraph("Анализ типичной ошибки", style="Title")
        _line(document, "Класс и предмет", f"{result.request.grade} класс, {result.request.subject_name}")
        _line(document, "Тема", result.request.topic)
        _line(document, "Описание ошибки", result.request.description)
        _line(document, "Вероятная картина", result.summary)
        document.add_heading("Гипотезы", level=1)
        confidence = {"low": "низкая", "medium": "средняя", "high": "высокая"}
        for index, hypothesis in enumerate(result.hypotheses, start=1):
            document.add_heading(f"Гипотеза {index}", level=2)
            _line(document, "Слой", ERROR_LAYER_LABELS[hypothesis.layer])
            if hypothesis.uud_group:
                _line(document, "Группа УУД", LABELS_RU[hypothesis.uud_group.value])
            _line(document, "Возможная причина", hypothesis.statement)
            _line(document, "Уверенность", confidence[hypothesis.confidence])
            if hypothesis.signs_to_check:
                document.add_paragraph("Что проверить")
                _bullets(document, hypothesis.signs_to_check)
        document.add_heading("Приёмы на следующий урок", level=1)
        for technique in result.techniques:
            document.add_heading(technique.name, level=2)
            document.add_paragraph(technique.how_to_apply)
            _line(document, "Ожидаемый сдвиг", technique.expected_shift)
        if result.quick_checks:
            document.add_heading("Мини-задания для проверки", level=1)
            for task in result.quick_checks:
                document.add_paragraph(task.student_text)
                _line(document, "Ожидаемый ответ", task.expected_answer)
                _line(document, "Что проверяет", task.what_it_checks)
        if result.teacher_reflection_question:
            _line(document, "Вопрос педагогу", result.teacher_reflection_question)
        document.add_heading("Ограничения", level=1)
        _bullets(document, result.limitations)
        return _bytes(document)
