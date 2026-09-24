"""Реализация KnowledgeBase — локальная нормативная база (G7).

ВЛАДЕЛЕЦ: Gemini (knowledge).
"""

from __future__ import annotations

import functools
import hashlib
import io
import json
import logging
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pymupdf
import yaml

from socrat.config import Settings
from socrat.contracts import (
    Chunk,
    ChunkKind,
    DocType,
    EduLevel,
    KnowledgeStatus,
    Outcome,
    OutcomeType,
    ProgressCallback,
    SearchHit,
    SourceDocument,
    Subject,
    TopicCheck,
    UpdateReport,
)
from socrat.knowledge.catalog import build_all_catalogs
from socrat.knowledge.chunker import chunk_document, chunk_pages_file, save_chunks
from socrat.knowledge.downloader import download_all, load_manifest, load_sources_registry
from socrat.knowledge.index import RUSSIAN_STOPWORDS, KnowledgeIndex, build_embeddings_for_index
from socrat.knowledge.parser import get_page_text, parse_and_save_pdf
from socrat.knowledge.retriever import Retriever

logger = logging.getLogger(__name__)

SK01_CASE_IDS = ("P01", "P02", "C01", "R01", "K01", "L01")

RUSSIAN_ENDINGS: tuple[str, ...] = (
    "евшимися",
    "овавшимися",
    "евшими",
    "овавшими",
    "ившимися",
    "ывшимися",
    "ившими",
    "ывшими",
    "ившихся",
    "ывшихся",
    "евшихся",
    "ующихся",
    "ющихся",
    "ающихся",
    "яющихся",
    "ившийся",
    "ывшийся",
    "евшийся",
    "ующийся",
    "ющийся",
    "ающийся",
    "яющийся",
    "анными",
    "янными",
    "енными",
    "енного",
    "енному",
    "анного",
    "янного",
    "анному",
    "янному",
    "анном",
    "янном",
    "енном",
    "анных",
    "янных",
    "енных",
    "анная",
    "янная",
    "енная",
    "анную",
    "янную",
    "енную",
    "анное",
    "янное",
    "енное",
    "анные",
    "янные",
    "енные",
    "тельными",
    "тельного",
    "тельному",
    "тельном",
    "тельных",
    "тельная",
    "тельную",
    "тельное",
    "тельные",
    "тельный",
    "ческими",
    "ческого",
    "ческому",
    "ческом",
    "ческих",
    "ческая",
    "ческую",
    "ческое",
    "ческие",
    "ческий",
    "ованными",
    "ованного",
    "ованному",
    "ованном",
    "ованных",
    "ованная",
    "ованную",
    "ованное",
    "ованные",
    "ованный",
    "ениями",
    "ениях",
    "ением",
    "ениям",
    "ений",
    "ение",
    "ения",
    "ению",
    "ении",
    "аниями",
    "аниях",
    "анием",
    "аниям",
    "аний",
    "ание",
    "ания",
    "анию",
    "ании",
    "остями",
    "остях",
    "остям",
    "остью",
    "ости",
    "ость",
    "ыми",
    "ими",
    "ого",
    "его",
    "ому",
    "ему",
    "ых",
    "их",
    "ую",
    "юю",
    "ая",
    "яя",
    "ое",
    "ее",
    "ые",
    "ие",
    "ым",
    "им",
    "ом",
    "ем",
    "ый",
    "ий",
    "ой",
    "ями",
    "ами",
    "ей",
    "ев",
    "ов",
    "ам",
    "ям",
    "ах",
    "ях",
    "ся",
    "сь",
    "ть",
    "ти",
    "те",
    "ет",
    "ут",
    "ют",
    "ит",
    "ат",
    "ят",
    "ил",
    "ыл",
    "ла",
    "ло",
    "ли",
    "е",
    "и",
    "ы",
    "а",
    "я",
    "о",
    "у",
    "ю",
    "ь",
)


@functools.lru_cache(maxsize=1)
def _morph():
    try:
        import pymorphy3

        return pymorphy3.MorphAnalyzer()
    except Exception:  # pragma: no cover - нет словарей
        return None


@functools.lru_cache(maxsize=50_000)
def stem_ru(word: str) -> str:
    """Основа слова: сначала лемма pymorphy3 («имена» → «имя»), затем отсечение окончания."""
    w = word.lower().replace("ё", "е")
    morph = _morph()
    if morph is not None and w.isalpha():
        w = morph.parse(w)[0].normal_form.replace("ё", "е")
    for end in RUSSIAN_ENDINGS:
        if w.endswith(end) and len(w) - len(end) >= 3:
            return w[: -len(end)]
    return w


def topic_coverage(topic: str, text: str) -> float:
    """Вычисляет долю значимых слов темы, найденных в тексте."""
    q_words = [
        w for w in re.findall(r"[а-яёa-z0-9]+", topic.lower()) if w not in RUSSIAN_STOPWORDS and len(w) >= 2
    ]
    if not q_words:
        q_words = [topic.lower()]

    cand_words = set(re.findall(r"[а-яёa-z0-9]+", text.lower()))
    cand_stems = {stem_ru(cw) for cw in cand_words}
    matched = 0
    for qw in q_words:
        qst = stem_ru(qw)
        if qst in cand_stems or any(len(qst) >= 4 and qst in cw for cw in cand_words):
            matched += 1
    return matched / len(q_words)


class LocalKnowledgeBase:
    """Реализует socrat.contracts.KnowledgeBase на локальных файлах."""

    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self._sources: dict[str, SourceDocument] = {}
        self._subjects_def: list[dict[str, Any]] = []
        self._outcomes: list[Outcome] = []
        self._outcomes_by_id: dict[str, Outcome] = {}
        self._problems: list[str] = []
        self._ready: bool = False
        self._last_update: datetime | None = None

        self.index: KnowledgeIndex = KnowledgeIndex([])
        self.retriever: Retriever = Retriever(self.index, self._sources, self.settings)

        self._load()

    def _load(self) -> None:
        """Загружает манифест, каталог, список предметов и поисковый индекс."""
        self._problems.clear()

        # 1. Загрузка subjects.yaml
        subjects_file = Path(__file__).resolve().parent / "subjects.yaml"
        if subjects_file.exists():
            try:
                raw_subs = yaml.safe_load(subjects_file.read_text(encoding="utf-8")) or {}
                self._subjects_def = sorted(raw_subs.get("subjects", []), key=lambda s: s.get("order", 99))
            except Exception as e:
                self._problems.append(f"Не удалось загрузить subjects.yaml: {e}")
        else:
            self._problems.append("Файл subjects.yaml не найден")

        # 2. Загрузка источников из manifest.json (или fallback на sources.yaml)
        manifest_path = self.settings.knowledge_dir / "manifest.json"
        manifest_docs = load_manifest(manifest_path)
        if manifest_docs:
            for sid, doc_data in manifest_docs.items():
                try:
                    self._sources[sid] = SourceDocument(
                        source_id=doc_data.get("source_id", sid),
                        title=doc_data.get("title", sid),
                        url=doc_data.get("url", ""),
                        doc_type=DocType(doc_data.get("doc_type", "frp")),
                        edu_levels=[EduLevel(lvl) for lvl in doc_data.get("edu_levels", [])],
                        subject_id=doc_data.get("subject_id"),
                        grades=doc_data.get("grades", []),
                        is_normative=doc_data.get("is_normative", True),
                        sha256=doc_data.get("sha256"),
                        pages=doc_data.get("pages"),
                        local_path=doc_data.get("local_path"),
                        note=doc_data.get("note"),
                    )
                except Exception as e:
                    logger.warning("Error parsing document %s from manifest: %s", sid, e)
        else:
            # Fallback: читаем sources.yaml
            src_list = load_sources_registry(self.settings.sources_file)
            for s in src_list:
                sid = s["source_id"]
                try:
                    self._sources[sid] = SourceDocument(
                        source_id=sid,
                        title=s.get("title", sid),
                        url=s.get("url", ""),
                        doc_type=DocType(s.get("doc_type", "frp")),
                        edu_levels=[EduLevel(lvl) for lvl in s.get("edu_levels", [])],
                        subject_id=s.get("subject_id"),
                        grades=s.get("grades", []),
                        is_normative=s.get("is_normative", True),
                        note=s.get("note"),
                    )
                except Exception as e:
                    logger.warning("Error creating fallback source %s: %s", sid, e)

        if not self._sources:
            self._problems.append("Нормативные источники не зарегистрированы")

        # 3. Загрузка каталога планируемых результатов
        catalog_dir = self.settings.knowledge_dir / "catalog"
        self._outcomes.clear()
        self._outcomes_by_id.clear()

        if catalog_dir.exists():
            for c_file in sorted(catalog_dir.glob("*.json")):
                try:
                    data = json.loads(c_file.read_text(encoding="utf-8"))
                    if isinstance(data, list):
                        for item in data:
                            out = Outcome.model_validate(item)
                            if out.outcome_id not in self._outcomes_by_id:
                                self._outcomes.append(out)
                                self._outcomes_by_id[out.outcome_id] = out
                except Exception as e:
                    self._problems.append(f"Ошибка чтения файла каталога {c_file.name}: {e}")

        if not self._outcomes:
            self._problems.append("Каталог планируемых результатов пуст или не найден")

        # 4. Загрузка поискового индекса
        index_dir = self.settings.knowledge_dir / "index"
        pages_dir = self.settings.knowledge_dir / "pages"
        all_chunks: list[Chunk] = []

        if (index_dir / "chunks.jsonl").exists():
            try:
                self.index = KnowledgeIndex.load(index_dir)
            except Exception as e:
                logger.warning("Failed to load KnowledgeIndex from %s: %s", index_dir, e)
                self.index = KnowledgeIndex([])
        else:
            # Если индекса на диске нет, проверяем сначала chunks/, затем pages/
            chunks_dir = self.settings.knowledge_dir / "chunks"
            if chunks_dir.exists():
                for ch_file in chunks_dir.glob("*.jsonl"):
                    with open(ch_file, encoding="utf-8") as f:
                        for line in f:
                            if line.strip():
                                all_chunks.append(Chunk.model_validate_json(line))

            if not all_chunks and pages_dir.exists():
                for pages_file in sorted(pages_dir.glob("*.jsonl")):
                    sid = pages_file.stem
                    doc = self._sources.get(sid)
                    s_id = doc.subject_id if doc else None
                    try:
                        doc_chunks = chunk_pages_file(pages_file, sid, s_id)
                        all_chunks.extend(doc_chunks)
                    except Exception as e:
                        logger.warning("Failed to chunk %s: %s", pages_file, e)

            self.index = KnowledgeIndex(all_chunks)
            if all_chunks:
                try:
                    self.index.save(index_dir)
                except Exception as e:
                    logger.debug("Could not save index to %s: %s", index_dir, e)

        # Проверяем, все ли страницы из pages/ присутствуют в индексе
        if pages_dir.exists():
            indexed_sources = {c.source_id for c in self.index.chunks}
            missing_pages = [pf for pf in sorted(pages_dir.glob("*.jsonl")) if pf.stem not in indexed_sources]
            if missing_pages:
                new_chunks: list[Chunk] = []
                for pf in missing_pages:
                    sid = pf.stem
                    doc = self._sources.get(sid)
                    s_id = doc.subject_id if doc else None
                    try:
                        doc_chunks = chunk_pages_file(pf, sid, s_id)
                        new_chunks.extend(doc_chunks)
                    except Exception as e:
                        logger.warning("Failed to chunk %s: %s", pf, e)
                if new_chunks:
                    all_chunks = list(self.index.chunks) + new_chunks
                    self.index = KnowledgeIndex(all_chunks)
                    try:
                        self.index.save(index_dir)
                    except Exception as e:
                        logger.debug("Could not save updated index to %s: %s", index_dir, e)

        self.retriever = Retriever(self.index, self._sources, self.settings)

        # 5. Диагностика целостности для админки (G13)
        for doc in self._sources.values():
            if doc.is_normative and doc.doc_type == DocType.FRP:
                p_file = pages_dir / f"{doc.source_id}.jsonl"
                if not p_file.exists():
                    self._problems.append(f"нет pages/ для {doc.source_id} — выполните build")
                has_outcomes = any(o.source_id == doc.source_id for o in self._outcomes)
                if not has_outcomes:
                    self._problems.append(f"нет результатов в каталоге для {doc.source_id} — выполните build")

        # База готова, если есть зарегистрированные источники и загруженный каталог
        self._ready = bool(self._sources and self._outcomes)

    def list_subjects(self, grade: int) -> list[Subject]:
        """Возвращает только те предметы для класса, по которым есть хотя бы один результат в каталоге (G15)."""
        try:
            EduLevel.for_grade(grade)
        except ValueError:
            return []

        # Только предметы, у которых для этого класса есть хотя бы один результат в каталоге
        catalog_subjects = {o.subject_id for o in self._outcomes if o.grade == grade and o.subject_id}

        res: list[Subject] = []
        for s in self._subjects_def:
            sid = s["subject_id"]
            if grade in s.get("grades", []) and sid in catalog_subjects:
                res.append(
                    Subject(
                        subject_id=sid,
                        name=s["name"],
                        grades=s.get("grades", []),
                    )
                )
        return res

    def get_subject(self, subject_id: str) -> Subject | None:
        for s in self._subjects_def:
            if s["subject_id"] == subject_id:
                return Subject(
                    subject_id=s["subject_id"],
                    name=s["name"],
                    grades=s.get("grades", []),
                )
        return None

    def check_topic(self, grade: int, subject_id: str, topic: str) -> TopicCheck:
        """Проверяет соответствие темы учителя программе класса по предмету (G10, G14, G17)."""
        try:
            EduLevel.for_grade(grade)
        except ValueError:
            return TopicCheck(
                in_program=False,
                confidence=0.0,
                matched_topics=[],
                suggestions=[],
                evidence=[],
            )

        valid_subjects = {s.subject_id for s in self.list_subjects(grade)}
        if not valid_subjects or subject_id not in valid_subjects:
            return TopicCheck(
                in_program=False,
                confidence=0.0,
                matched_topics=[],
                suggestions=[],
                evidence=[],
            )

        target_outcomes = [
            o
            for o in self._outcomes
            if o.grade == grade and (o.subject_id is None or o.subject_id == subject_id)
        ]

        if not topic or not topic.strip():
            return TopicCheck(
                in_program=False,
                confidence=0.0,
                matched_topics=[],
                suggestions=self._extract_suggestions(grade, subject_id, target_outcomes),
                evidence=[],
            )

        # 1. Поиск по фрагментам содержания и предметных результатов индекса (если индекс не пуст)
        hits: list[SearchHit] = []
        if len(self.index.chunks) > 0:
            raw_hits = self.retriever.search(
                query=topic,
                grade=grade,
                subject_id=subject_id,
                kinds=[ChunkKind.CONTENT, ChunkKind.SUBJECT_RESULT],
                k=5,
            )
            # Фильтруем результаты: тема должна покрывать >= 60% значимых слов (G16)
            hits = [h for h in raw_hits if topic_coverage(topic, h.chunk.text) >= 0.60]

        # 2. Поиск по формулировкам результатов каталога этого класса и предмета (G14, G17)
        outcome_hits: list[SearchHit] = []
        if target_outcomes:
            temp_chunks = [
                Chunk(
                    chunk_id=o.outcome_id,
                    source_id=o.source_id,
                    page=o.page,
                    text=o.text,
                    kind=ChunkKind.SUBJECT_RESULT,
                    grade=o.grade,
                    subject_id=o.subject_id,
                    section=o.section,
                )
                for o in target_outcomes
            ]
            temp_index = KnowledgeIndex(temp_chunks)
            raw_outcome_hits = Retriever(temp_index, self._sources).search(
                topic,
                grade=grade,
                subject_id=subject_id,
                k=5,
            )
            outcome_hits = [h for h in raw_outcome_hits if topic_coverage(topic, h.chunk.text) >= 0.60]

        # Проверяем, есть ли тема в результатах этого класса из каталога (G17)
        in_this_grade_outcome = any(topic_coverage(topic, o.text) >= 0.60 for o in target_outcomes) or any(
            h.score >= 0.4 for h in outcome_hits
        )

        # Проверяем, есть ли тема в результатах других классов этого же предмета (G17)
        other_grade_outcomes = [
            o
            for o in self._outcomes
            if o.subject_id == subject_id and o.grade is not None and o.grade != grade
        ]
        in_other_grade_outcome = any(topic_coverage(topic, o.text) >= 0.60 for o in other_grade_outcomes)

        has_explicit_grade_chunk = any(h.chunk.grade == grade and h.score >= 0.4 for h in hits)

        in_this_grade = in_this_grade_outcome or (has_explicit_grade_chunk and not in_other_grade_outcome)

        matched_topics: list[str] = []
        best_score = hits[0].score if hits else 0.0

        if in_this_grade:
            in_prog = True
            outcome_best_score = outcome_hits[0].score if outcome_hits else 0.0
            best_score = max(best_score, outcome_best_score)
            if not hits and outcome_hits:
                hits = outcome_hits
            for h in hits:
                if h.score >= 0.4:
                    first_line = h.chunk.text.split("\n")[0].strip()
                    matched_topics.append(first_line[:80])
        else:
            in_prog = False
            has_general_chunk_hit = any(h.score >= 0.4 for h in hits)
            if has_general_chunk_hit or in_other_grade_outcome:
                matched_topics = ["Тема есть в программе предмета, но в другом классе"]
                best_score = min(0.5, max(best_score, 0.4))
            else:
                matched_topics = []
                best_score = 0.1

        # 3. Подсказки тем (до 5 коротких названий)
        suggestions: list[str] = []
        if not in_prog:
            suggestions = self._extract_suggestions(grade, subject_id, target_outcomes)

        return TopicCheck(
            in_program=in_prog,
            confidence=min(1.0, max(0.1, best_score)),
            matched_topics=list(dict.fromkeys(matched_topics))[:3],
            suggestions=suggestions[:5],
            evidence=hits,
        )

    def _extract_suggestions(self, grade: int, subject_id: str, target_outcomes: list[Outcome]) -> list[str]:
        """Извлекает до 5 подсказок тем программы для указанного класса и предмета."""
        suggestions: list[str] = []
        # Сначала пытаемся извлечь темы из фрагментов раздела «Содержание обучения» этого класса
        content_chunks = [
            c
            for c in self.index.chunks
            if c.grade == grade
            and (c.subject_id is None or c.subject_id == subject_id)
            and c.kind == ChunkKind.CONTENT
        ]
        seen_topics: set[str] = set()
        ignore_keywords = {"универсальные", "планирование", "результаты", "деятельность", "действия"}
        for ch in content_chunks:
            for line in ch.text.split("\n"):
                for sentence in re.split(r"[;\.]", line):
                    s_clean = sentence.strip()
                    if 10 <= len(s_clean) <= 60 and not any(w in s_clean.lower() for w in ignore_keywords):
                        s_title = s_clean[0].upper() + s_clean[1:]
                        if s_title not in seen_topics:
                            seen_topics.add(s_title)
                            suggestions.append(s_title)
                            if len(suggestions) >= 5:
                                break
                if len(suggestions) >= 5:
                    break

        # Если подсказок меньше 5 (или нет content_chunks), берём из каталога результатов (G14)
        if len(suggestions) < 5 and target_outcomes:
            for o in target_outcomes:
                first_phrase = o.text.split("\n")[0].split(";")[0].split(".")[0].strip()
                clean_phrase = re.sub(
                    r"^(Числа и вычисления|Алгебраические выражения|Уравнения и неравенства|Функции|Наглядная геометрия|Геометрические фигуры|Язык и речь|СИСТЕМА ЯЗЫКА|Текст|Фонетика|Орфография|Лексикология|Морфемика|Морфология|Синтаксис)\s+",
                    "",
                    first_phrase,
                    flags=re.I,
                ).strip()
                if 10 <= len(clean_phrase) <= 60 and not any(
                    w in clean_phrase.lower() for w in ignore_keywords
                ):
                    s_title = clean_phrase[0].upper() + clean_phrase[1:]
                    if s_title not in seen_topics:
                        seen_topics.add(s_title)
                        suggestions.append(s_title)
                        if len(suggestions) >= 5:
                            break

        # Fallback для 3 класса математики при необходимости
        if len(suggestions) < 5 and grade == 3 and subject_id == "math":
            math3_default = [
                "Умножение и деление в пределах 100",
                "Решение текстовых задач в одно-два действия",
                "Сложение и вычитание в пределах 1000",
                "Периметр и площадь прямоугольника",
                "Деление с остатком",
            ]
            for d in math3_default:
                if d not in suggestions:
                    suggestions.append(d)
                    if len(suggestions) >= 5:
                        break

        return suggestions[:5]

    def get_outcomes(
        self,
        grade: int,
        subject_id: str,
        types: list[OutcomeType] | None = None,
        query: str | None = None,
        k: int = 30,
    ) -> list[Outcome]:
        """Возвращает планируемые результаты для класса и предмета.

        Для math/3 обязательно первыми включает P01..L01 из кейса SK01.
        """
        try:
            EduLevel.for_grade(grade)
        except ValueError:
            return []

        # 1. Предметные результаты этого класса и предмета
        subject_outcomes = [
            o
            for o in self._outcomes
            if o.type == OutcomeType.SUBJECT
            and o.grade == grade
            and o.subject_id == subject_id
            and (not types or o.type in types)
        ]

        # 2. Метапредметные и личностные результаты уровня
        meta_types = (
            OutcomeType.COGNITIVE,
            OutcomeType.COMMUNICATIVE,
            OutcomeType.REGULATORY,
            OutcomeType.PERSONAL,
        )
        meta_outcomes = [
            o
            for o in self._outcomes
            if o.type in meta_types
            and (o.subject_id is None or o.subject_id == subject_id)
            and (not types or o.type in types)
        ]

        sk01_items: list[Outcome] = []
        regular_items: list[Outcome] = []

        combined = subject_outcomes + meta_outcomes
        seen_ids: set[str] = set()

        # Для math/3 кейса SK01: всегда первыми
        if grade == 3 and subject_id == "math":
            for cid in SK01_CASE_IDS:
                o = self._outcomes_by_id.get(cid)
                if o and (not types or o.type in types):
                    sk01_items.append(o)
                    seen_ids.add(o.outcome_id)

        for o in combined:
            if o.outcome_id not in seen_ids:
                regular_items.append(o)
                seen_ids.add(o.outcome_id)

        # Ранжирование по поисковому запросу при наличии
        if query:
            q_lower = query.lower()
            q_words = set(q_lower.split())

            def _rank_key(item: Outcome) -> float:
                i_words = set(item.text.lower().split())
                return len(q_words & i_words)

            regular_items.sort(key=_rank_key, reverse=True)

        res = sk01_items + regular_items
        return res[:k]

    def get_outcome(self, outcome_id: str) -> Outcome | None:
        return self._outcomes_by_id.get(outcome_id)

    def search(
        self,
        query: str,
        grade: int | None = None,
        subject_id: str | None = None,
        kinds: list[ChunkKind] | None = None,
        k: int = 8,
        include_school_materials: bool = True,
    ) -> list[SearchHit]:
        """Гибридный поиск по фрагментам документов (BM25 + эмбеддинги)."""
        return self.retriever.search(
            query=query,
            grade=grade,
            subject_id=subject_id,
            kinds=kinds,
            k=k,
            include_school_materials=include_school_materials,
        )

    def get_page_text(self, source_id: str, page: int) -> str | None:
        """Возвращает очищенный текст страницы PDF."""
        pages_dir = self.settings.knowledge_dir / "pages"
        return get_page_text(source_id, page, pages_dir)

    def get_source(self, source_id: str) -> SourceDocument | None:
        return self._sources.get(source_id)

    def status(self) -> KnowledgeStatus:
        return KnowledgeStatus(
            documents=list(self._sources.values()),
            chunks_total=len(self.index.chunks),
            outcomes_total=len(self._outcomes),
            embeddings_enabled=bool(self.settings.embeddings_base_url and self.settings.embeddings_model),
            last_update=self._last_update,
            ready=self._ready,
            problems=list(self._problems),
        )

    async def update(self, progress: ProgressCallback | None = None) -> UpdateReport:
        """Полный цикл обновления: G2 (скачивание) -> G3 (парсинг) -> G4 (нарезка) -> G5 (каталог) -> G6 (индекс)."""
        start_time = datetime.now(UTC)

        # 1. G2: Скачивание
        if progress:
            await progress("Скачивание нормативных документов…")
        report = await download_all(self.settings, progress)

        # 2. G3: Парсинг страниц PDF
        raw_dir = self.settings.knowledge_dir / "raw"
        pages_dir = self.settings.knowledge_dir / "pages"
        pages_dir.mkdir(parents=True, exist_ok=True)

        for pdf_path in raw_dir.glob("*.pdf"):
            sid = pdf_path.stem
            if progress:
                await progress(f"Парсинг PDF {sid}…")
            parse_and_save_pdf(pdf_path, sid, pages_dir, sources_file=self.settings.sources_file)

        # 3. G4: Нарезка на фрагменты (чанкинг)
        chunks_dir = self.settings.knowledge_dir / "chunks"
        chunks_dir.mkdir(parents=True, exist_ok=True)

        all_chunks: list[Chunk] = []
        for pages_file in pages_dir.glob("*.jsonl"):
            sid = pages_file.stem
            doc = self._sources.get(sid)
            s_id = doc.subject_id if doc else None
            if progress:
                await progress(f"Нарезка разделов {sid}…")
            doc_chunks = chunk_pages_file(pages_file, sid, s_id)
            save_chunks(doc_chunks, chunks_dir / f"{sid}.jsonl")
            all_chunks.extend(doc_chunks)

        # 4. G5: Сборка каталогов результатов
        if progress:
            await progress("Сборка каталогов планируемых результатов…")
        build_all_catalogs(self.settings)

        # 5. G6: Сборка поискового индекса и эмбеддингов
        if progress:
            await progress("Построение поискового индекса BM25…")
        index_dir = self.settings.knowledge_dir / "index"
        self.index = KnowledgeIndex(all_chunks)
        self.index.save(index_dir)

        if self.settings.embeddings_base_url and self.settings.embeddings_model:
            if progress:
                await progress("Генерация векторных эмбеддингов…")
            await build_embeddings_for_index(self.index, self.settings, index_dir)

        # Перезагружаем состояние
        self._load()
        self._last_update = datetime.now(UTC)

        report.chunks_total = len(self.index.chunks)
        report.outcomes_total = len(self._outcomes)
        report.duration_s = (datetime.now(UTC) - start_time).total_seconds()

        if progress:
            await progress("Обновление нормативной базы завершено.")
        return report

    async def ingest_school_material(
        self,
        filename: str,
        content: bytes,
        progress: ProgressCallback | None = None,
    ) -> SourceDocument:
        """Загрузка PDF/DOCX школы как школьного материала (is_normative=False)."""
        sha = hashlib.sha256(content).hexdigest()
        source_id = f"SCHOOL-{sha[:8]}"

        if progress:
            await progress(f"Обработка школьного материала {filename}…")

        pages: list[dict[str, Any]] = []

        if filename.lower().endswith(".docx"):
            # Извлечение из DOCX блоками по 3000 символов (G7)
            import docx

            doc = docx.Document(io.BytesIO(content))
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            page_size = 3000
            if not full_text:
                pages.append({"page": 1, "text": ""})
            else:
                for idx, start_idx in enumerate(range(0, len(full_text), page_size), start=1):
                    pages.append({"page": idx, "text": full_text[start_idx : start_idx + page_size]})
        else:
            # Извлечение из PDF с помощью pymupdf
            try:
                doc_pdf = pymupdf.open(stream=content, filetype="pdf")
                for idx, page in enumerate(doc_pdf, start=1):
                    pages.append({"page": idx, "text": page.get_text("text")})
            except Exception as e:
                logger.warning("Could not parse PDF content via pymupdf: %s", e)
                pages.append({"page": 1, "text": content.decode("utf-8", errors="ignore")})

        source_doc = SourceDocument(
            source_id=source_id,
            title=filename,
            url=f"local://{filename}",
            doc_type=DocType.SCHOOL_MATERIAL,
            is_normative=False,
            sha256=sha,
            pages=len(pages),
        )
        self._sources[source_id] = source_doc

        # Нарезаем и индексируем
        material_chunks = chunk_document(source_id, pages, subject_id=None)
        for ch in material_chunks:
            self.index.add_chunk(ch)

        # Сохраняем обновленный индекс
        index_dir = self.settings.knowledge_dir / "index"
        self.index.save(index_dir)
        self.retriever = Retriever(self.index, self._sources, self.settings)

        if progress:
            await progress(f"Школьный материал {filename} успешно проиндексирован.")

        return source_doc
