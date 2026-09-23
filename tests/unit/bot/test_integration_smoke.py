"""Full fake-mode service path for the first demo scenario."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from socrat.app import build_services, close_services
from socrat.bot.auth import authenticate_password
from socrat.config import get_settings
from socrat.contracts import (
    ErrorAnalysisRequest,
    Feedback,
    GenerationRequest,
    LevelChoice,
    Rating,
    Role,
    UsageRecord,
)
from socrat.storage import SessionStore, session_factory


async def test_first_demo_scenario_across_services(tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("APP_SECRET", "test-integration-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "test-admin-password")
    monkeypatch.setenv("DATABASE_URL", f"sqlite+aiosqlite:///{tmp_path / 'integration.db'}")
    monkeypatch.setenv("USE_FAKES", "true")
    monkeypatch.setenv("LLM_PRICE_IN_PER_1M", "0")
    monkeypatch.setenv("LLM_PRICE_OUT_PER_1M", "0")
    get_settings.cache_clear()
    settings = get_settings()
    services = await build_services(settings)
    try:
        assert (
            await authenticate_password("test-admin-password", services.teachers, settings)
        ).role == Role.ADMIN
        created = await services.teachers.create("Педагог", Role.TEACHER)
        teacher = await authenticate_password(created.plain_password, services.teachers, settings)
        assert teacher is not None and teacher.teacher_id == created.account.teacher_id

        sessions = SessionStore(session_factory(services._storage_engine), settings)
        await sessions.create(12345, teacher)
        assert (await sessions.get(12345)).nick == "Педагог"

        subjects = services.knowledge.list_subjects(3)
        subject = next(item for item in subjects if item.subject_id == "math")
        topic = "Умножение и деление в пределах 100"
        assert services.knowledge.check_topic(3, subject.subject_id, topic).in_program
        request = GenerationRequest(
            grade=3,
            subject_id=subject.subject_id,
            subject_name=subject.name,
            topic=topic,
            level=LevelChoice.ALL,
            task_count=4,
            author_role=teacher.role,
        )
        assert not (await services.generator.check_request(request)).blocked
        work = await services.generator.generate(request)
        assert len(work.variants) == 3
        assert all(len(variant.tasks) == 4 for variant in work.variants)
        await services.works.save(work, teacher.teacher_id)
        assert (await services.works.get(work.work_id)).work_id == work.work_id

        student = Document(BytesIO(services.exporter.student_docx(work)))
        teacher_doc = Document(BytesIO(services.exporter.teacher_docx(work)))
        assert any("Вариант А" in paragraph.text for paragraph in student.paragraphs)
        assert any("Методические материалы" in paragraph.text for paragraph in teacher_doc.paragraphs)

        task = work.variants[0].tasks[0]
        await services.feedback.add(
            Feedback(
                work_id=work.work_id,
                task_id=task.task_id,
                rating=Rating.UP,
                teacher_id=teacher.teacher_id,
            )
        )
        assert (await services.feedback.top_examples(3, "math"))[0].task_id == task.task_id
        assert services.observer.observe(task, task.expected_answer).task_id == task.task_id

        analysis_request = ErrorAnalysisRequest(
            grade=3,
            subject_id="math",
            subject_name=subject.name,
            topic=topic,
            description="В заданиях на деление дети иногда выбирают умножение.",
            author_role=teacher.role,
        )
        assert not (await services.analyzer.check(analysis_request)).blocked
        analysis = await services.analyzer.analyze(analysis_request)
        assert analysis.hypotheses
        assert Document(BytesIO(services.exporter.analysis_docx(analysis))).paragraphs

        await services.usage.add(
            UsageRecord(kind="generation", model=work.meta.model, teacher_id=teacher.teacher_id)
        )
        await services.usage.add(
            UsageRecord(kind="analysis", model=analysis.meta.model, teacher_id=teacher.teacher_id)
        )
        assert (await services.usage.summary(30)).requests == 2
        assert services.knowledge.status().ready
        await sessions.delete(12345)
        assert await sessions.get(12345) is None
    finally:
        await close_services(services)
        get_settings.cache_clear()
