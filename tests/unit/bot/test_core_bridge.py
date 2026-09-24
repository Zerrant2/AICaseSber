"""Exercise the real core through the app while knowledge is still a fixture."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

import socrat.knowledge
from socrat.app import build_services, close_services
from socrat.config import Settings
from socrat.contracts import ErrorAnalysisRequest, GenerationRequest, LevelChoice
from socrat.testing.fakes import FakeKnowledgeBase
from socrat.testing.scripted import ScriptedLLM


async def test_real_core_with_fixture_knowledge_and_storage(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delattr(socrat.knowledge, "LocalKnowledgeBase", raising=False)
    llm = ScriptedLLM()
    monkeypatch.setattr("socrat.llm.OpenAICompatibleClient", lambda settings: llm)
    settings = Settings(
        _env_file=None,
        app_secret="core-bridge-test-secret",
        database_url=f"sqlite+aiosqlite:///{tmp_path / 'core_bridge.db'}",
        use_fakes=False,
    )
    services = await build_services(settings)
    try:
        assert isinstance(services.knowledge, FakeKnowledgeBase)
        request = GenerationRequest(
            grade=3,
            subject_id="math",
            subject_name="Математика",
            topic="Умножение и деление в пределах 100",
            level=LevelChoice.ALL,
            task_count=4,
        )
        assert not (await services.generator.check_request(request)).blocked
        work = await services.generator.generate(request)
        assert len(work.variants) == 3
        assert all(len(variant.tasks) == 4 for variant in work.variants)
        assert len(llm.calls) == 3

        await services.works.save(work, None)
        assert (await services.works.get(work.work_id)).work_id == work.work_id

        student = Document(BytesIO(services.exporter.student_docx(work)))
        teacher = Document(BytesIO(services.exporter.teacher_docx(work)))
        assert any("Вариант А" in paragraph.text for paragraph in student.paragraphs)
        assert any("Методические материалы" in paragraph.text for paragraph in teacher.paragraphs)

        task = work.variants[0].tasks[0]
        assert services.observer.observe(task, task.expected_answer).task_id == task.task_id
        analysis = await services.analyzer.analyze(
            ErrorAnalysisRequest(
                grade=3,
                subject_id="math",
                subject_name="Математика",
                topic="Площадь и периметр прямоугольника",
                description="Дети перемножают стороны, когда нужно найти периметр, и складывают для площади.",
            )
        )
        assert analysis.hypotheses
        assert Document(BytesIO(services.exporter.analysis_docx(analysis))).paragraphs
        assert len(llm.calls) == 4
    finally:
        await close_services(services)
