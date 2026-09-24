"""Check the Word files built from the SK01 contract fixtures."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from socrat.config import get_settings
from socrat.contracts import DiagnosticWork, ErrorAnalysisResult, Exporter
from socrat.export import DocxExporter

ROOT = Path(__file__).resolve().parents[3]


@pytest.fixture(autouse=True)
def settings(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("LLM_PRICE_IN_PER_1M", "0")
    monkeypatch.setenv("LLM_PRICE_OUT_PER_1M", "0")
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def fixture_work() -> DiagnosticWork:
    return DiagnosticWork.model_validate_json(
        (ROOT / "data" / "fixtures" / "work_math3_all.json").read_text(encoding="utf-8")
    )


def document_text(document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_student_docx_has_only_student_fields() -> None:
    work = fixture_work()
    exporter = DocxExporter()
    assert isinstance(exporter, Exporter)
    document = Document(BytesIO(exporter.student_docx(work)))
    text = document_text(document)
    assert text.count("Имя ______________") == 3
    assert all(label in text for label in ("Вариант А", "Вариант Б", "Вариант В"))
    assert "Ответ:" not in text
    for forbidden in ("Лёгкий", "Базовый", "Сложный"):
        assert forbidden not in text
    for variant in work.variants:
        for task in variant.tasks:
            for step in task.solution_steps:
                assert step.text not in text
    assert document.styles["Normal"].font.name == "Times New Roman"
    assert document.styles["Normal"].font.size == Pt(12)
    assert abs(document.sections[0].left_margin - Mm(20)) < 650
    assert abs(document.sections[0].page_width - Mm(210)) < 650
    assert document.element.body.xml.count('w:type="page"') == 2


def test_teacher_docx_has_outcomes_and_pages() -> None:
    work = fixture_work()
    document = Document(BytesIO(DocxExporter().teacher_docx(work)))
    text = document_text(document)
    assert "Методические материалы для учителя" in text
    assert "Коды результатов — внутренние идентификаторы системы" in text
    for variant in work.variants:
        for task in variant.tasks:
            assert f"Ответ: {task.expected_answer}" in text
            for link in task.outcome_links:
                assert link.outcome_id in text
                assert any(
                    link.outcome_id in row.cells[1].text and str(link.page) == row.cells[5].text
                    for table in document.tables
                    for row in table.rows
                    if len(row.cells) == 8
                )
    assert all(
        row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        for table in document.tables
        if len(table.columns) == 8
        for row in table.rows[1:]
    )
    assert all(
        paragraph.paragraph_format.keep_with_next
        for table in document.tables
        if len(table.rows) > 1
        for cell in table.rows[0].cells
        for paragraph in cell.paragraphs
    )
    assert all(
        table.columns[0].width >= Mm(40)
        and all(
            paragraph.paragraph_format.keep_with_next
            for row in table.rows[:-1]
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
        for table in document.tables
        if len(table.columns) == 8
    )
    map_headings = [
        index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == "Карта результатов"
    ]
    assert map_headings and all(
        'w:type="page"' in document.paragraphs[index - 1]._p.xml for index in map_headings
    )
    assert "Числовой ответ сам по себе не подтверждает" in text
    assert "Личностная направленность: Личностная направленность" not in text
    assert abs(document.sections[0].page_width - Mm(297)) < 650


def test_analysis_docx_and_safe_filenames() -> None:
    result = ErrorAnalysisResult.model_validate_json(
        (ROOT / "data" / "fixtures" / "error_analysis_math3.json").read_text(encoding="utf-8")
    )
    document = Document(BytesIO(DocxExporter().analysis_docx(result)))
    text = document_text(document)
    assert result.request.description in text
    assert result.summary in text
    assert all(hypothesis.statement in text for hypothesis in result.hypotheses)
    assert all(technique.name in text for technique in result.techniques)
    work = fixture_work()
    work.request.topic = "Умножение: / деление? *"
    names = DocxExporter().filenames(work)
    assert names[0].startswith("Задания_3кл_")
    assert names[1].startswith("Методичка_3кл_")
    assert not any(char in name for name in names for char in '<>:"/\\|?*')
