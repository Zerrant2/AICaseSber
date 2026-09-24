"""12 проверок кейса SK01 (офлайн: сценарная LLM + справочник кейса).

Проверяется логика системы: чертёж, проверки кодом, ограничители, выгрузка, разбор ответов.
Прогон на реальной модели — tests/case_sk01/test_case_llm.py (маркер llm).
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path

import pytest
from docx import Document
from pydantic import ValidationError

from socrat.contracts import (
    DiagnosticWork,
    GenerationRequest,
    GuardrailCode,
    Level,
    LevelChoice,
    OutcomeType,
    Severity,
    SubjectObservation,
    UUDGroup,
    UUDObservation,
    student_view,
)
from socrat.core.mathcheck import evaluate, fmt, parse_number
from socrat.core.observation import RuleResponseObserver
from socrat.export.docx_exporter import DocxExporter

ROOT = Path(__file__).resolve().parents[2]
REF = ROOT / "data" / "reference" / "sk01"

CASE_IDS = {"P01", "P02", "C01", "R01", "K01", "L01"}


def _req(teacher_request: dict, **kw) -> GenerationRequest:
    base = dict(
        grade=teacher_request["grade"],
        subject_id="math",
        subject_name=teacher_request["subject"],
        topic=teacher_request["topic"],
        level=LevelChoice.ALL,
        task_count=teacher_request["tasks_per_variant"],
    )
    base.update(kw)
    return GenerationRequest(**base)


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


async def test_T01_three_versions(core, teacher_request, record):
    gen, _, _ = core
    work = await gen.generate(_req(teacher_request))
    assert [v.level for v in work.variants] == [Level.EASY, Level.BASIC, Level.ADVANCED]
    for v in work.variants:
        assert len(v.tasks) == 4 and v.reflection.questions
        assert v.time_plan.within_range
        assert 15 <= v.time_plan.total_minutes <= 20
    record(
        "T01",
        "; ".join(
            f"{v.student_label}: {len(v.tasks)} задания + {len(v.reflection.questions)} вопр. "
            f"рефлексии, {v.time_plan.total_minutes:g} мин"
            for v in work.variants
        ),
    )


def test_T02_example_arithmetic_and_student_export(record):
    variants = json.loads((REF / "example_variants.json").read_text(encoding="utf-8"))
    checked = 0
    for v in variants:
        for t in v["tasks"]:
            exprs = [e for e in re.findall(r"\d+\s*[×:+−\-]\s*\d+", t["solution_example"])]
            values = {evaluate(e) for e in exprs}
            assert parse_number(str(t["expected_answer"])) in values, t["task_id"]
            checked += 1
    assert checked == 12
    work = DiagnosticWork.model_validate_json(
        (ROOT / "data/fixtures/work_math3_all.json").read_text(encoding="utf-8")
    )
    text = _docx_text(DocxExporter().student_docx(work))
    for v in work.variants:
        for t in v.tasks:
            for s in t.solution_steps:
                assert s.text not in text
    assert "Лёгкий" not in text and "Сложный" not in text
    record(
        "T02",
        f"Все {checked} ответов примера пересчитаны вычислителем кода и совпали; в ученическом .docx "
        "нет ни одного шага решения и ярлыков уровня",
    )


async def test_T03_outcome_map(core, teacher_request, record):
    gen, _, _ = core
    work = await gen.generate(_req(teacher_request))
    n = 0
    for v in work.variants:
        for t in v.tasks:
            types = {lk.type for lk in t.outcome_links}
            assert OutcomeType.SUBJECT in types and t.subject_goal
            assert t.uud_indicators and t.personal_orientation
            assert OutcomeType.PERSONAL in types
            for lk in t.outcome_links:
                assert lk.outcome_id in CASE_IDS and lk.section and lk.page >= 1 and lk.source_id
            n += 1
    assert any("не номера пунктов" in x for x in work.limitations)
    record(
        "T03",
        f"{n} карт: предметный результат, наблюдаемые УУД (с триггером), личностная направленность, "
        "источник FRP-MATH-2025 с разделом и страницей; в ограничениях — «коды не номера пунктов ФГОС»",
    )


@pytest.mark.xfail(
    reason="Визуальная/практическая подача отложена по решению заказчика (23.09); "
    "модель Task.modes и student_view уже поддерживают подмену формулировки",
    strict=False,
)
def test_T04_presentation_modes(record):
    work = DiagnosticWork.model_validate_json(
        (ROOT / "data/fixtures/work_math3_all.json").read_text(encoding="utf-8")
    )
    record("T04", "Отложено: генерация визуальной/практической подачи не реализована в MVP")
    assert all(t.modes for v in work.variants for t in v.tasks)


async def test_T05_unknown_norm_item(core, teacher_request, record):
    gen, _, _ = core
    r = await gen.check_request(_req(teacher_request, topic="Умножение и деление по пункту 42.7.3 ФГОС НОО"))
    issue = next(i for i in r.issues if i.code == GuardrailCode.UNKNOWN_NORM_REFERENCE)
    assert "нет в фиксированном справочнике" in issue.message_ru and issue.suggestions
    assert all(s.split(":")[0] in CASE_IDS for s in issue.suggestions)
    record("T05", issue.message_ru[:160] + "… Предложено: " + "; ".join(issue.suggestions[:2]))


def test_T06_number_without_explanation(record):
    responses = json.loads((REF / "synthetic_responses.json").read_text(encoding="utf-8"))
    work = DiagnosticWork.model_validate_json(
        (ROOT / "data/fixtures/work_math3_all.json").read_text(encoding="utf-8")
    )
    ids = {"EASY": "A", "BASIC": "B", "ADVANCED": "C"}
    obs = RuleResponseObserver()
    short = [r for r in responses if r["response_id"].endswith("SHORT")]
    for r in short:
        lvl, num = r["task_id"].split("-")
        o = obs.observe(work.find_task(f"{ids[lvl]}-{num}"), r["response"])
        assert o.subject_status == SubjectObservation.CORRECT
        assert all(i.status != UUDObservation.OBSERVED for i in o.uud)
    record(
        "T06",
        f"{len(short)} ответов «только число»: предметно — верно; по УУД — «недостаточно наблюдений» "
        f"(пример: {o.summary_ru[:120]}…)",
    )


def test_T07_reflection_difficulty(record):
    r = RuleResponseObserver().observe_reflection(
        "Что было трудно?", "Мне было трудно, я не понял третье задание"
    )
    assert "не оценка личности" in r.neutral_note and "диагноз" in r.neutral_note
    record("T07", r.neutral_note)


async def test_T08_eight_minutes(core, teacher_request, record):
    gen, _, _ = core
    r = await gen.check_request(_req(teacher_request, requested_minutes=8))
    issue = next(i for i in r.issues if i.code == GuardrailCode.TIME_OUT_OF_RANGE)
    assert "15–20" in issue.message_ru and issue.suggestions
    record("T08", issue.message_ru)


async def test_T09_levels_differ_by_support(core, teacher_request, record):
    gen, _, _ = core
    work = await gen.generate(_req(teacher_request))
    easy, basic, adv = work.variants
    assert all(t.support for t in easy.tasks)
    assert not any(t.support for t in basic.tasks)
    groups_adv = [{i.group for i in t.uud_indicators} for t in adv.tasks]
    assert sum(UUDGroup.COMMUNICATIVE in g for g in groups_adv) >= 2
    assert adv.time_plan.total_minutes <= 20
    assert not any("только в объёме" in w.message_ru for w in work.warnings)
    record(
        "T09",
        f"Лёгкий: опоры у {sum(bool(t.support) for t in easy.tasks)}/4; базовый: без подсказок; сложный: "
        f"объяснение в {sum(UUDGroup.COMMUNICATIVE in g for g in groups_adv)}/4, типы заданий "
        f"{[t.kind.value for t in adv.tasks]}; время {adv.time_plan.total_minutes:g} мин",
    )


async def test_T10_unknown_subject_or_grade(core, teacher_request, record):
    gen, _, _ = core
    r = await gen.check_request(_req(teacher_request, subject_id="chemistry", subject_name="Химия"))
    assert (
        r.blocked
        and r.issues[0].code == GuardrailCode.SUBJECT_UNKNOWN
        and "Математика" in r.issues[0].suggestions
    )
    with pytest.raises(ValidationError):
        _req(teacher_request, grade=12)
    record("T10", r.issues[0].message_ru + " Класс 12 отклоняется на уровне запроса (1–11).")


async def test_T11_change_numbers_recalculated(core, teacher_request, record):
    gen, _, _ = core
    work = await gen.generate(_req(teacher_request, level=LevelChoice.BASIC))
    before = work.variants[0].tasks[3]
    new = await gen.regenerate_task(work, Level.BASIC, 4, wish="замени 4 на 5, 6 на 7")
    t = new.variants[0].tasks[3]
    assert t.expected_answer == "30" and t.checks.math_verified and t.level == Level.BASIC
    for s in t.solution_steps:
        assert evaluate(s.expression) == parse_number(s.result)
    record(
        "T11",
        f"«{before.student_text[:60]}…» → «{t.student_text[:60]}…»; ответ {before.expected_answer} → "
        f"{t.expected_answer}; шаги: " + ", ".join(f"{s.expression}={s.result}" for s in t.solution_steps),
    )


async def test_T12_full_fgos_guarantee(core, teacher_request, record):
    gen, _, _ = core
    req = _req(teacher_request, teacher_note="Гарантируй полное соответствие всему ФГОС")
    r = await gen.check_request(req)
    issue = next(i for i in r.issues if i.code == GuardrailCode.FULL_FGOS_GUARANTEE)
    assert issue.severity == Severity.WARN and "не может охватить весь ФГОС" in issue.message_ru
    work = await gen.generate(req)
    assert any("не всему ФГОС" in x for x in work.limitations)
    record("T12", issue.message_ru)


def test_student_view_projection_is_clean():
    work = DiagnosticWork.model_validate_json(
        (ROOT / "data/fixtures/work_math3_all.json").read_text(encoding="utf-8")
    )
    dumped = student_view(work).model_dump_json()
    assert "solution" not in dumped and "expected_answer" not in dumped
    assert fmt(evaluate("3*4")) == "12"
